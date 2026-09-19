#!/usr/bin/env python3
"""Apply deterministic publication/advisor Word styles to Pandoc DOCX exports.

The presets implement the repository's document handoff needs while following
the bundled document-skill token contract. This script changes formatting only;
it does not alter the Markdown source or scientific content.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


PRESETS = {
    "narrative_proposal": {
        "font": "Calibri", "body_after": 8, "line": 1.333,
        "h1": (16, 18, 10), "h2": (13, 12, 6), "h3": (12, 8, 4),
        "table_fill": "F4F6F9", "table_font": 8.5,
    },
    "standard_business_brief": {
        "font": "Calibri", "body_after": 6, "line": 1.10,
        "h1": (16, 16, 8), "h2": (13, 12, 6), "h3": (12, 8, 4),
        "table_fill": "F2F4F7", "table_font": 9.0,
    },
    "compact_reference_guide": {
        "font": "Calibri", "body_after": 6, "line": 1.25,
        "h1": (16, 18, 10), "h2": (13, 14, 7), "h3": (12, 10, 5),
        "table_fill": "E8EEF5", "table_font": 9.0,
    },
    "decision_memo": {
        "font": "Arial", "body_after": 6, "line": 1.10,
        "h1": (16, 12, 6), "h2": (13, 10, 5), "h3": (12, 8, 4),
        "table_fill": "F2F4F7", "table_font": 9.0,
    },
}

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x66, 0x66, 0x66)


def set_font(run, name: str, size: float | None = None, *, color=None,
             bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style(style, *, font: str, size: float, color, before: float,
              after: float, line: float, bold: bool = False,
              justify: bool = False) -> None:
    style.font.name = font
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if justify:
        fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def set_cell_margins(cell, *, top=80, bottom=80, start=120, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom),
                        ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def column_widths(table, total=9360) -> list[int]:
    cols = len(table.columns)
    if cols == 1:
        return [total]
    scores = []
    for col in range(cols):
        lengths = [len(row.cells[col].text.strip()) for row in table.rows]
        scores.append(max(8.0, min(48.0, (sum(lengths) / max(len(lengths), 1)) + 5)))
    minimum = 820 if cols >= 5 else 1050
    available = total - minimum * cols
    score_total = sum(scores)
    widths = [minimum + int(available * score / score_total) for score in scores]
    widths[-1] += total - sum(widths)
    return widths


def style_table(table, preset: dict, total_width: int = 9360) -> None:
    widths = column_widths(table, total=total_width)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag, attrs in (
        ("w:tblW", {"w:w": str(total_width), "w:type": "dxa"}),
        ("w:tblInd", {"w:w": "120", "w:type": "dxa"}),
        ("w:tblLayout", {"w:type": "fixed"}),
    ):
        node = tbl_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tbl_pr.append(node)
        for key, value in attrs.items():
            node.set(qn(key), value)

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "3")
        node.set(qn("w:color"), "D7DEE8")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row_index, row in enumerate(table.rows):
        row_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)
        if row_index == 0:
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            row_pr.append(tbl_header)
        for col_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[col_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                tc_pr = cell._tc.get_or_add_tcPr()
                shd = tc_pr.find(qn("w:shd"))
                if shd is None:
                    shd = OxmlElement("w:shd")
                    tc_pr.append(shd)
                shd.set(qn("w:fill"), preset["table_fill"])
            short_column = max(
                len(r.cells[col_index].text.strip()) for r in table.rows
            ) <= 22
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if short_column
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    set_font(
                        run, preset["font"],
                        min(preset["table_font"], 8.0)
                        if total_width > 9360 else preset["table_font"],
                        color=INK,
                        bold=True if row_index == 0 else None,
                    )


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def style_document(path: Path, preset_name: str, label: str | None,
                   landscape: bool = False) -> None:
    preset = PRESETS[preset_name]
    doc = Document(path)
    for section in doc.sections:
        section.page_width = Inches(11 if landscape else 8.5)
        section.page_height = Inches(8.5 if landscape else 11)
        margin = 0.5 if landscape else 1.0
        section.top_margin = Inches(margin)
        section.right_margin = Inches(margin)
        section.bottom_margin = Inches(margin)
        section.left_margin = Inches(margin)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header
        hp = header.paragraphs[0]
        hp.text = label or "Ottoman QML OCR | Publication Evidence"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.paragraph_format.space_after = Pt(0)
        for run in hp.runs:
            set_font(run, preset["font"], 8.5, color=MUTED)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Page "
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_field(fp, " PAGE ")
        for run in fp.runs:
            set_font(run, preset["font"], 8.5, color=MUTED)

    justify = preset_name == "narrative_proposal"
    for style_name in ("Normal", "Body Text", "First Paragraph"):
        if style_name in doc.styles:
            set_style(
                doc.styles[style_name], font=preset["font"], size=11,
                color=RGBColor(0, 0, 0), before=0,
                after=preset["body_after"], line=preset["line"],
                justify=justify,
            )
    if "Compact" in doc.styles:
        set_style(
            doc.styles["Compact"], font=preset["font"], size=11,
            color=RGBColor(0, 0, 0), before=0, after=4,
            line=1.25 if preset_name == "compact_reference_guide" else 1.167,
        )

    for name, token, color in (
        ("Heading 1", preset["h1"], BLUE),
        ("Heading 2", preset["h2"], BLUE),
        ("Heading 3", preset["h3"], DARK_BLUE),
    ):
        if name in doc.styles:
            size, before, after = token
            set_style(
                doc.styles[name], font=preset["font"], size=size,
                color=color, before=before, after=after, line=1.0, bold=True,
            )
            doc.styles[name].paragraph_format.keep_with_next = True

    nonempty = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip()]
    if nonempty:
        title = nonempty[0]
        title.paragraph_format.space_before = Pt(0)
        title.paragraph_format.space_after = Pt(14)
        title.paragraph_format.line_spacing = 1.0
        title.paragraph_format.keep_with_next = True
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title_size = 19 if len(title.text) > 70 else 23
        for run in title.runs:
            set_font(run, preset["font"], title_size, color=INK, bold=True)
        if len(nonempty) > 1 and nonempty[1].style.name not in {
            "Heading 1", "Heading 2", "Heading 3"
        }:
            meta = nonempty[1]
            meta.paragraph_format.space_after = Pt(12)
            for run in meta.runs:
                set_font(run, preset["font"], 10, color=MUTED, italic=True)

    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.style.name == "Source Code":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_together = True
            for run in paragraph.runs:
                set_font(run, "Courier New", 8.5, color=INK)

    for table in doc.tables:
        style_table(table, preset, total_width=14400 if landscape else 9360)

    previous_was_table = False
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            previous_was_table = True
            continue
        if child.tag == qn("w:p"):
            if previous_was_table:
                Paragraph(child, doc).paragraph_format.space_before = Pt(8)
            previous_was_table = False

    for shape in doc.inline_shapes:
        if shape.width > Inches(6.4):
            scale = Inches(6.4) / shape.width
            shape.width = int(shape.width * scale)
            shape.height = int(shape.height * scale)

    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--preset", choices=sorted(PRESETS), required=True)
    parser.add_argument("--label", default=None)
    parser.add_argument("--landscape", action="store_true")
    args = parser.parse_args()
    for path in args.files:
        style_document(path, args.preset, args.label, landscape=args.landscape)
        print(f"Styled {path} with {args.preset}")


if __name__ == "__main__":
    main()
